from flask import Flask, redirect, render_template, request, jsonify, url_for
from openpyxl import load_workbook
import pandas as pd
from graphapiutils import calcular_valores_coluna
from models import Mesa, db
import os

from models import Pesquisa

# Carregar a planilha na inicialização
downloads_folder = os.path.join(os.getcwd(), 'downloads')
file_path = os.path.join(downloads_folder, 'Tabela_Preco.xlsx')

dataframe = pd.read_excel(file_path, sheet_name='Planilha6')
print(dataframe)

dataframe['ESTADO'] = dataframe['ESTADO'].str.strip()  # Remover espaços extras
dataframe['MUNICÍPIO'] = dataframe['MUNICÍPIO'].str.strip()
print("Colunas no DataFrame:", dataframe.columns.tolist())


df = dataframe.dropna(subset=['ESTADO', 'MUNICÍPIO'])
df.columns = df.columns.str.strip().str.upper().str.replace(' ', '_')

#lógica para captar informações dos usuários
def get_estados():
    """Retorna a lista de estados únicos com validação."""
    estados = df['ESTADO'].unique().tolist()
    if not estados:
        return []
    return estados

def get_municipios(estado):
    """Retorna a lista de municípios únicos com base no estado selecionado com validação."""
    municipios = df[df['ESTADO'] == estado]['MUNICÍPIO'].unique().tolist()
    if not municipios:
        return []
    return municipios

def get_regioes():
    """Retorna a lista de regiões únicas com validação."""
    regioes = df['REGIÃO'].unique().tolist()
    if not regioes:
        return []
    return regioes

#inicia-se o framework e o banco de dados
app = Flask(__name__)
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///database.db'
app.config.from_object('config.Config')
db.init_app(app)

#rotas
@app.route('/filtrar_regioes', methods=['POST'])
def filtrar_regioes():
    estado = request.json.get('estado')
    municipio = request.json.get('municipio')
    regioes_filtradas = []

    if estado and municipio:
        # Filtrar as regiões com base no estado e município
        regioes_filtradas = df[(df['ESTADO'] == estado) & (df['MUNICÍPIO'] == municipio)]['REGIÃO'].unique().tolist()

        if regioes_filtradas:
            try:
                # Criar uma nova entrada no banco de dados
                nova_pesquisa = Pesquisa(
                    estado=estado,
                    municipio=municipio,
                    regiao=int(regioes_filtradas[0])  # Seleciona a primeira região como exemplo
                )
                db.session.add(nova_pesquisa)
                db.session.commit()
                print(f"Pesquisa salva: {nova_pesquisa}")
            except Exception as e:
                db.session.rollback()
                print(f"Erro ao salvar no banco de dados: {e}")

    return jsonify({'regioes': regioes_filtradas})


@app.route('/', methods=['GET', 'POST'])
def index():
    estados = get_estados()
    municipios = []
    regioes = []

    if request.method == 'POST':
        estado = request.form.get('estado')
        municipio = request.form.get('municipio')
        regiao = request.form.get('regiao')

        if estado:
            municipios = get_municipios(estado)

        if estado and municipio:
            # Filtrar as regiões com base no estado e município
            regioes = df[(df['ESTADO'] == estado) & (df['MUNICÍPIO'] == municipio)]['REGIÃO'].unique().tolist()

        if estado and municipio and regiao:
            try:
                regiao_int = int(regiao)
                dados_filtrados = df[(df['ESTADO'] == estado) & (df['MUNICÍPIO'] == municipio) & (df['REGIÃO'] == regiao_int)]

                if not dados_filtrados.empty:
                    # Gravar no banco de dados usando o modelo Pesquisa
                    try:
                        nova_pesquisa = Pesquisa(
                            estado=estado,
                            municipio=municipio,
                            regiao=regiao_int
                        )
                        db.session.add(nova_pesquisa)
                        db.session.commit()
                    except Exception as e:
                        mensagem_erro = f"Erro ao salvar no banco de dados: {str(e)}"
                        return render_template('index.html', estados=estados, municipios=municipios, regioes=regioes, mensagem_erro=mensagem_erro)

                    return redirect(url_for('mesas', regiao=regiao, estado=estado, municipio=municipio))
                else:
                    mensagem_erro = "Nenhum dado encontrado para os critérios selecionados."
                    return render_template('index.html', estados=estados, municipios=municipios, regioes=regioes, mensagem_erro=mensagem_erro)

            except ValueError:
                mensagem_erro = "Selecione uma região válida."
                return render_template('index.html', estados=estados, municipios=municipios, regioes=regioes, mensagem_erro=mensagem_erro)

        mensagem_erro = "Por favor, selecione o estado, município e região."
        return render_template('index.html', estados=estados, municipios=municipios, regioes=regioes, mensagem_erro=mensagem_erro)

    return render_template('index.html', estados=estados, municipios=municipios, regioes=regioes)


@app.route('/municipios', methods=['GET'])
def get_municipios_ajax():
    estado = request.args.get('estado')
    if estado:
        municipios = get_municipios(estado)
        return jsonify(municipios)
    return jsonify([])


@app.route("/mesas", methods=["GET", "POST"])
def calculadora_mesas():
    # Caminho para o arquivo no diretório "downloads"
    downloads_folder = os.path.join(os.getcwd(), 'downloads')
    file_path = os.path.join(downloads_folder, 'PRECO-MODULO-JULLYEN.xlsx')

    wb = load_workbook(file_path, data_only=True)
    all_sheets_data = []
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        data = sheet.values
        columns = next(data)
        df_sheet = pd.DataFrame(data, columns=columns)
        all_sheets_data.append(df_sheet)

    # Concatenar os dados de todas as abas em um único DataFrame
    df = pd.concat(all_sheets_data, ignore_index=True)

    # Padronizar os nomes das colunas
    df.columns = df.columns.str.strip().str.lower().str.replace(' ', '_')
    print("Colunas no DataFrame após padronização:", df.columns.tolist())

    # Verificar se as colunas necessárias existem
    required_columns = ['poste', 'fileiras', 'graus', 'região', 'módulo', 'preço_da_mesa_/módulo']
    if not all(col in df.columns for col in required_columns):
        return f"Erro: Colunas necessárias não encontradas no arquivo. Colunas atuais: {df.columns.tolist()}"

    # Padronizar os dados apenas nas colunas relevantes
    cols_to_standardize = ['poste', 'fileiras', 'graus', 'região', 'módulo']
    for col in cols_to_standardize:
        if col in df.columns:
            df[col] = df[col].astype(str).str.strip().str.lower()

    # Extrair opções únicas para os selects
    postes = df['poste'].dropna().unique().tolist()
    fileiras = df['fileiras'].dropna().unique().tolist()
    graus = df['graus'].dropna().unique().tolist()
    regioes = df['região'].dropna().unique().tolist()
    modulos = df['módulo'].dropna().unique().tolist()

    preco = None
    total = None
    mensagem_erro = None
    regiao_selecionada = None

    if request.method == "POST":
        # Captura os valores selecionados no formulário
        poste = request.form.get("poste", "").strip().lower() if request.form.get("poste") else None
        fileira = request.form.get("fileiras", "").strip().lower() if request.form.get("fileiras") else None
        grau = request.form.get("graus", "").strip().lower() if request.form.get("graus") else None
        regiao_selecionada = request.form.get("regiao", "").strip().lower() if request.form.get("regiao") else None
        modulo = request.form.get("modulo", "").strip().lower() if request.form.get("modulo") else None
        quantidade = request.form.get("quantidade")

        # Validar se todos os campos necessários foram preenchidos
        if not all([poste, fileira, grau, regioes, modulo]):
            mensagem_erro = "Todos os campos devem ser preenchidos."
            return render_template(
                "result.html",
                postes=postes, fileiras=fileiras, graus=graus, regioes=regioes, modulos=modulos,
                preco=preco, total=total, mensagem_erro=mensagem_erro
            )


        try:
            # Filtrar dados no DataFrame
            resultado = df[
                (df['poste'] == poste) &
                (df['fileiras'] == fileira) &
                (df['graus'] == grau) &
                #(df['região'] == regiao) &
                (df['módulo'] == modulo)
            ]

            # Verificar se há resultados
            if not resultado.empty:
                preco = float(resultado['preço_da_mesa_/módulo'].iloc[0])  # Obtém o preço unitário
                if quantidade and quantidade.isdigit():
                    total = preco * int(quantidade)  # Calcula o total

                    #salvando no BD
                    nova_mesa = Mesa(
                        poste=poste,
                        fileiras=fileira,
                        graus=grau,
                        regiao=regiao_selecionada,
                        modulo=modulo,
                        preco=preco
                    )
                    db.session.add(nova_mesa)
                    db.session.commit()
                else:
                    mensagem_erro = "Quantidade inválida. Digite um número maior que zero."
            else:
                mensagem_erro = "Nenhum resultado encontrado para os critérios selecionados."

        except Exception as e:
            mensagem_erro = f"Erro ao calcular: {str(e)}"

    return render_template(
        "result.html",
        postes=postes, fileiras=fileiras, graus=graus, regioes=regioes, modulos=modulos,
        preco=preco, total=total, mensagem_erro=mensagem_erro, regiao_selecionada=regiao_selecionada
    )




if __name__ == '__main__':
    with app.app_context():
        db.create_all()  # Cria as tabelas no banco de dados
        print("Banco de dados inicializado com sucesso!")
        app.run(debug=True)
