import logging
import requests
from msal import ConfidentialClientApplication
from decouple import config
import os
import glob
from openpyxl import load_workbook


# Configuração do logging
logging.basicConfig(
    level=logging.DEBUG,  # Define o nível de logging
    format="%(asctime)s - %(levelname)s - %(message)s",  # Formato do log
    datefmt="%Y-%m-%d %H:%M:%S"  # Formato da data
)


# Configurações
CLIENT_ID = config("CLIENT_ID")
CLIENT_SECRET = config("CLIENT_SECRET")
TENANT_ID = config("TENANT_ID")
DRIVE_ID = config("DRIVEID")
DIRECTORY_PATH = "Comercial/CALCULADORA"

# Endpoint da API Graph
GRAPH_API_URL = f"https://graph.microsoft.com/v1.0/drives/{DRIVE_ID}/root:/{DIRECTORY_PATH}:/children"

# Função para autenticação e obtenção do token para acessar a API
def get_access_token():
    app = ConfidentialClientApplication(
        CLIENT_ID,
        authority=f"https://login.microsoftonline.com/{TENANT_ID}",
        client_credential=CLIENT_SECRET
    )
    result = app.acquire_token_for_client(scopes=["https://graph.microsoft.com/.default"])
    if "access_token" in result:
        return result["access_token"]
    else:
        raise Exception(f"Erro ao obter token: {result}")


#para pegar todas as pastas que estão nas pastas da microsoft
def list_files(access_token):
    logging.info("Iniciando a função list_files.")
    
    headers = {"Authorization": f"Bearer {access_token}"}
    logging.debug(f"Headers preparados para a requisição: {headers}")
    logging.debug(f"Endpoint utilizado: {GRAPH_API_URL}")
    
    try:
        response = requests.get(GRAPH_API_URL, headers=headers)
        logging.debug(f"Resposta recebida. Status code: {response.status_code}")
        logging.debug(f"Conteúdo da resposta: {response.text[:500]}")  # Mostra os primeiros 500 caracteres da resposta
        
        if response.status_code == 200:
            logging.info("Arquivos listados com sucesso.")
            return response.json()
        else:
            logging.error(f"Erro ao listar arquivos. Status code: {response.status_code}, Resposta: {response.text}")
            raise Exception(f"Erro ao listar arquivos: {response.status_code}, {response.text}")
    except Exception as e:
        logging.critical(f"Exceção capturada: {e}")
        raise



# Função para criar uma pasta e verificar se o arquivo já existe antes do download
def download_file(access_token, download_url, directory, file_name):
    # Cria a pasta se ela não existir
    if not os.path.exists(directory):
        os.makedirs(directory)
        print(f"Pasta {directory} criada com sucesso.")
    # Caminho completo do arquivo
    file_path = os.path.join(directory, file_name)
    # Verifica se o arquivo já existe
    if os.path.exists(file_path):
        print(f"Arquivo {file_name} já existe. Pulando download.")
        return
    # Faz o download do arquivo
    headers = {"Authorization": f"Bearer {access_token}"}
    response = requests.get(download_url, headers=headers, stream=True)
    if response.status_code == 200:
        with open(file_path, "wb") as file:
            for chunk in response.iter_content(chunk_size=8192):
                file.write(chunk)
        print(f"Arquivo {file_name} baixado com sucesso em {directory}.")
    else:
        raise Exception(f"Erro ao baixar arquivo: {response.status_code}, {response.text}")

# Fluxo principal
try:
    token = get_access_token()
    files = list_files(token)
    download_directory = "downloads"  # Diretório onde os arquivos serão salvos
    for file in files.get("value", []):
        if "@microsoft.graph.downloadUrl" in file:
            file_name = file["name"]
            download_url = file["@microsoft.graph.downloadUrl"]
            download_file(token, download_url, download_directory, file_name)
except Exception as e:
    print(f"Erro: {e}")

#encontra o ultimo arquivo salvo na pasta.
def get_latest_file(folder, pattern):
    """Encontra o arquivo mais recente no diretório que corresponde ao padrão."""
    files = glob.glob(os.path.join(folder, pattern))   # Procura arquivos que correspondem ao padrão
    if not files:
        raise FileNotFoundError("Nenhum arquivo correspondente encontrado na pasta downloads.")
    latest_file = max(files, key=os.path.getmtime)  # Seleciona o arquivo mais recente
    # Verifica se o arquivo é um arquivo Excel
    if not latest_file.lower().endswith('.xlsx'):
        raise ValueError("O arquivo encontrado não é um arquivo Excel (.xlsx).")
    print(f"Arquivo encontrado: {latest_file}")  # Para depuração
    return latest_file


def calcular_valores_coluna(download_folder, file_name):
    # Construir o caminho do arquivo
    file_path = os.path.join(download_folder, file_name)

    # Carregar a planilha com os valores calculados
    wb = load_workbook(file_path, data_only=True)  # 'data_only=True' pega valores calculados, não fórmulas
    sheet = wb.active  # Seleciona a primeira aba ativa

    # Iterar sobre a coluna ou célula com fórmulas
    valores = []
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=19, max_col=19):  # Ajuste os índices da coluna
        for cell in row:
            valores.append(cell.value)  # Adiciona os valores calculados à lista

    # Exibe os valores calculados
    print("Valores calculados:")
    for valor in valores:
        if valor is not None:  # Verifica se o valor não é None
            print(f"{valor:.2f}")
        else:
            print("None")

    # Soma os valores, se necessário
    total = sum(filter(None, valores))  # Filtra valores None antes de somar
    print(f"Total da coluna: {total:.2f}")

    return valores, total
   